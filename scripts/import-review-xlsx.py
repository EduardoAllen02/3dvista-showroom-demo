#!/usr/bin/env python3
"""
Reads the review workbook filled by the client (scripts/build-review-xlsx.py) and writes:

  tour-project/febal-casa/product-facts/review.reviewed.json
      {validated_fact_ids, rejected_fact_ids, promoted_palette_groups, excluded_palette_groups, reviewer, at}
      read by scripts/build-catalog-v2.ts (`--strict` then uses only validated facts)
  tour-project/febal-casa/product-facts/review-corrections.json
      every row with a NO, a correction or a note, per sheet, to apply by hand
      (synonyms to add/remove, color families, harmonies signed, answers to open questions…)

Rules:
  Piezas / Modelos: ¿OK? = SI validates the row's facts; NO rejects them (the assistant then
    says the data is not confirmed) and the correction is exported.
  Opciones oficiales / Paletas genéricas: a collection (one fact) is validated when at least one
    of its rows is SI and none is NO; a NO keeps it pending and exports the correction (a single
    wrong color must not wipe the whole collection).
  Paletas por confirmar: SI turns that line palette into real options of the listed models (offered
    without caveat); the models named in "Modelos donde NO aplica" and a NO stop mentioning it at all.
    Blank keeps today's behavior: mentioned with a "check its page" caveat.

    python scripts/import-review-xlsx.py <workbook.xlsx> [--reviewer "Andrea"]
"""
import datetime as dt
import json
import os
import re
import sys
import unicodedata
from collections import defaultdict

from openpyxl import load_workbook

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
FACTS = os.path.join(ROOT, "tour-project", "febal-casa", "product-facts")
OK_COLS = ("¿OK?", "¿Se aprueba?", "¿Aplica de verdad a todos estos modelos?")
REVIEW_COLS = OK_COLS + ("Corrección", "Nota", "Agregar sinónimos", "Quitar sinónimos", "Familia correcta",
                         "Material correcto", "Respuesta", "Modelos donde NO aplica")


def rows(ws):
    head = [c.value for c in ws[1]]
    for r in ws.iter_rows(min_row=2, values_only=True):
        yield dict(zip(head, r))


def norm(v):
    return str(v).strip().upper() if v is not None else ""


def fold(s):
    return "".join(c for c in unicodedata.normalize("NFD", str(s).lower()) if unicodedata.category(c) != "Mn")


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    path = sys.argv[1]
    reviewer = sys.argv[sys.argv.index("--reviewer") + 1] if "--reviewer" in sys.argv else "Andrea"
    with open(os.path.join(ROOT, "clients", "febal-casa", "catalog.v2.json"), encoding="utf-8") as f:
        model_names = {m["id"]: m["name"] for m in json.load(f)["models"]}
    wb = load_workbook(path, read_only=True)
    validated, rejected, promoted, excluded_groups = set(), set(), set(), set()
    corrections = defaultdict(list)
    option_votes = defaultdict(lambda: {"SI": 0, "NO": 0})

    for ws in wb.worksheets:
        if ws.title == "Léeme":
            continue
        for row in rows(ws):
            ok = next((norm(row.get(k)) for k in OK_COLS if row.get(k) not in (None, "")), "")
            facts = [f for f in str(row.get("fact_ids") or "").split(",") if f]
            touched = ok == "NO" or any(row.get(k) not in (None, "") for k in REVIEW_COLS if k not in OK_COLS)
            if ws.title in ("Piezas (88)", "Modelos"):
                if ok == "SI":
                    validated.update(facts)
                elif ok == "NO":
                    rejected.update(facts)
            elif ws.title in ("Opciones oficiales", "Paletas genéricas"):
                for f in facts:
                    if ok in ("SI", "NO"):
                        option_votes[f][ok] += 1
            elif ws.title == "Paletas por confirmar":
                # "Arden, Couple (FEB-039)" -> {"arden", "couple"}: whole names or ids, never substrings.
                excluded = {t.strip() for t in re.split(r"[,;\n]", re.sub(r"\(.*?\)", "", fold(row.get("Modelos donde NO aplica") or "")))} - {""}
                for gid in filter(None, str(row.get("group_ids") or "").split(",")):
                    mid = gid.split("/")[0]
                    named_out = mid in excluded or fold(model_names.get(mid, "")) in excluded
                    if ok == "NO" or named_out:
                        excluded_groups.add(gid)
                    elif ok == "SI":
                        promoted.add(gid)
            if touched:
                corrections[ws.title].append({k: v for k, v in row.items() if v not in (None, "")})

    for f, v in option_votes.items():
        if v["SI"] and not v["NO"]:
            validated.add(f)
    validated -= rejected

    at = dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    review = {"validated_fact_ids": sorted(validated), "rejected_fact_ids": sorted(rejected),
              "promoted_palette_groups": sorted(promoted), "excluded_palette_groups": sorted(excluded_groups), "reviewer": reviewer, "at": at, "source": os.path.basename(path)}
    with open(os.path.join(FACTS, "review.reviewed.json"), "w", encoding="utf-8") as f:
        json.dump(review, f, ensure_ascii=False, indent=1)
    with open(os.path.join(FACTS, "review-corrections.json"), "w", encoding="utf-8") as f:
        json.dump({"reviewer": reviewer, "at": at, "sheets": corrections}, f, ensure_ascii=False, indent=1, default=str)
    print(f"validados {len(validated)} · rechazados {len(rejected)} · paletas promovidas {len(promoted)} grupos, descartadas {len(excluded_groups)} · "
          "filas con corrección/nota: " + ", ".join(f"{k} {len(v)}" for k, v in corrections.items()))
    print("Siguiente: aplicar review-corrections.json y reconstruir con  npx tsx scripts/build-catalog-v2.ts febal-casa --strict")


if __name__ == "__main__":
    main()

#!/usr/bin/env python3
"""
Builds the chatbot's database workbook: every piece, model, option, palette decision and
word the assistant uses, with a README in Spanish ("Léeme") and Italian ("Leggimi").
There are no review or question columns: open decisions live in the other workbook
(scripts/build-decisions-xlsx.py). Columns with a green "✎" header are editable; an edited
copy goes back through scripts/import-database-xlsx.py.

    python scripts/build-database-xlsx.py [out.xlsx]    # default: ~/Downloads/febal-casa-base-de-datos.xlsx
"""
import datetime as dt
import os
import re
import sys

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from PIL import Image

sys.stdout.reconfigure(encoding="utf-8")   # the Windows console defaults to cp1252
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import febal_database as db  # noqa: E402

TODAY = dt.date.today().isoformat()
DEFAULT_OUT = os.path.join(os.path.expanduser("~"), "Downloads", "febal-casa-base-de-datos.xlsx")
OUT = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_OUT
THUMBS = os.path.join(db.ROOT, ".scratch", "database-thumbs")
FORMAT_VERSION = "1"

RED, GREEN = "C8102E", "1E6B52"
HEAD_RO = PatternFill("solid", fgColor=RED)
HEAD_ED = PatternFill("solid", fgColor=GREEN)
HEAD_FONT = Font(bold=True, color="FFFFFF")
EDIT = PatternFill("solid", fgColor="EEF7F1")
NOTE_FILLS = {"CORREGIDO": PatternFill("solid", fgColor="DCE9F7"), "DUDA": PatternFill("solid", fgColor="FFE0B2"),
              "PENDIENTE": PatternFill("solid", fgColor="F8D7DA")}
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Border(bottom=Side(style="thin", color="DDDDDD"))
NAME_RE = re.compile(r"Andrea|\bEdd\b")


def thumb(pid):
    src = next((os.path.join(db.CAPTURES, f) for f in (f"{pid}-ui.png", f"{pid}.png") if os.path.exists(os.path.join(db.CAPTURES, f))), None)
    if not src:
        return None
    os.makedirs(THUMBS, exist_ok=True)
    dst = os.path.join(THUMBS, f"{pid}.jpg")
    with Image.open(src) as im:
        im = im.convert("RGB")
        im.thumbnail((300, 190))
        im.save(dst, "JPEG", quality=80)
    return XLImage(dst)


def write_sheet(wb, sh):
    ws = wb.create_sheet(sh.title)
    ws.append([c.header for c in sh.cols])
    for i, c in enumerate(sh.cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = c.width
        cell = ws.cell(1, i)
        cell.fill, cell.font, cell.alignment = HEAD_ED if c.editable else HEAD_RO, HEAD_FONT, Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 36
    for r, row in enumerate(sh.rows, start=2):
        ws.append(row.values)
        for i, c in enumerate(sh.cols, start=1):
            cell = ws.cell(r, i)
            cell.alignment, cell.border = WRAP, THIN
            if c.editable:
                cell.fill = EDIT
        for ci, url in row.links.items():
            if url:
                cell = ws.cell(r, ci + 1)
                cell.hyperlink, cell.font = url, Font(color="0563C1", underline="single")
        for ci, hexv in row.swatches.items():
            if re.fullmatch(r"#[0-9a-fA-F]{6}", hexv or ""):
                ws.cell(r, ci + 1).fill = PatternFill("solid", fgColor=hexv[1:].upper())
        for ci, kind in row.fills.items():
            if kind in NOTE_FILLS:
                ws.cell(r, ci + 1).fill = NOTE_FILLS[kind]
        if row.photo:
            col = next(i for i, c in enumerate(sh.cols, start=1) if c.kind == "photo")
            img = thumb(row.photo)
            if img:
                ws.add_image(img, f"{get_column_letter(col)}{r}")
        if row.height:
            ws.row_dimensions[r].height = row.height
    n = ws.max_row
    for i, c in enumerate(sh.cols, start=1):
        if c.choices:
            dv = DataValidation(type="list", formula1='"' + ",".join(c.choices) + '"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{get_column_letter(i)}2:{get_column_letter(i)}{n}")
    ws.column_dimensions[get_column_letter(len(sh.cols))].hidden = True   # "clave"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(sh.cols))}{n}"
    ws.freeze_panes = sh.freeze
    return ws


# ------------------------------------------------------------------ README (es / it)
class Readme:
    WIDTHS = [26, 40, 44, 40, 30]

    def __init__(self, ws):
        self.ws, self.r = ws, 1
        for i, w in enumerate(self.WIDTHS, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.sheet_view.showGridLines = False

    def _lines(self, text, width):
        return max(1, sum(-(-max(len(part), 1) // max(int(width * 0.95), 1)) for part in str(text).split("\n")))

    def para(self, text, font=None):
        ws = self.ws
        ws.merge_cells(start_row=self.r, start_column=1, end_row=self.r, end_column=len(self.WIDTHS))
        c = ws.cell(self.r, 1, text)
        c.alignment = WRAP
        if font:
            c.font = font
        ws.row_dimensions[self.r].height = 16 * self._lines(text, sum(self.WIDTHS)) + (10 if font and font.size and font.size > 12 else 3)
        self.r += 1

    def title(self, text):
        self.r += 1
        self.para(text, Font(bold=True, size=13, color=RED))

    def table(self, headers, rows, fills=None):
        ws, fills = self.ws, fills or {}
        for j, h in enumerate(headers, start=1):
            c = ws.cell(self.r, j, h)
            c.fill, c.font, c.alignment = HEAD_RO, HEAD_FONT, Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[self.r].height = 20
        self.r += 1
        for row in rows:
            for j, v in enumerate(row, start=1):
                c = ws.cell(self.r, j, v)
                c.alignment, c.border = WRAP, THIN
            if row[0] in fills:
                ws.cell(self.r, 1).fill = fills[row[0]]
            ws.row_dimensions[self.r].height = 16 * max(self._lines(v, self.WIDTHS[j]) for j, v in enumerate(row)) + 4
            self.r += 1


def stats(d, sheets):
    by = {s.title: s for s in sheets}
    pl = list(d.placements.values())
    yes = lambda k: sum(1 for p in pl if p[k]["confirmado"] == "SI")
    dec = [r.values[5] for r in by["Paletas por modelo"].rows]
    return {
        "pieces": len(pl), "color": yes("color"), "material": yes("material"), "forma": yes("forma"), "finish": yes("opcion_oficial"),
        "models": len(d.attrs), "options": len(by["Opciones oficiales"].rows), "pal_rows": len(dec),
        "aplica": dec.count("aplica"), "aviso": dec.count("aviso"), "no": dec.count("no aplica"),
        "pal_contents": len(by["Acabados de las paletas"].rows), "concepts": len(d.onto["concepts"]),
        "synonyms": sum(len(v) for c in d.onto["concepts"] for v in c.get("synonyms", {}).values()),
        "colors": len(by["Colores → familia"].rows), "collections": len(by["Materiales por colección"].rows),
        "relations": len(d.onto["relations"]), "corrections": len(by["Correcciones"].rows),
    }


def readme_es(ws, s):
    R = Readme(ws)
    R.para("Febal Casa · base de datos del asistente", Font(bold=True, size=16, color=RED))
    R.para(f"Generada el {TODAY}. Es exactamente lo que usa el asistente del tour para responder.")

    R.title("Qué es este archivo")
    R.para("Aquí está todo lo que sabe el asistente del showroom: cada pieza del tour, cada modelo de la web oficial, los acabados que puede ofrecer, "
           "las palabras que entiende y cómo las agrupa. Sirve para dos cosas: ver con exactitud qué sabe, y corregirlo o completarlo directamente aquí.")

    R.title("Cómo corregir algo")
    R.table(["Paso", "Qué hacer", "", "", ""], [
        ("1", "Busca la hoja y la fila. Los filtros de la primera fila ayudan.", "", "", ""),
        ("2", "Cambia solo las columnas con encabezado verde y ✎ (celdas verde claro).", "", "", ""),
        ("3", "Escribe como en las demás filas: familias de color y materiales con su nombre en español (p. ej. «verde», «lacado mate»), "
              "varios separados por coma. Los nombres válidos están en «Sinónimos», columna «Nombre (es)». En «Decisión» y «Estado» elige de la lista.", "", "", ""),
        ("4", "No agregues, borres ni reordenes filas, y no cambies las columnas de encabezado rojo: vienen de la web oficial o de las fotos del tour. "
              "Si ahí ves algo mal, avísalo al devolver el archivo.", "", "", ""),
        ("5", "Guarda y devuelve el archivo. Se revisa qué cambió, se aplica, se regenera la base del asistente, se prueba y recibes este archivo actualizado.", "", "", ""),
    ])

    R.title("Qué cambia en el asistente cuando editas")
    R.table(["Hoja", "Columna", "Efecto", "Ejemplo", ""], [
        ("Piezas (88)", "Color y material de la parte principal, forma vista", "Cómo describe la pieza expuesta y en qué búsquedas aparece", "FEB-036 Vivienne: material → «Nabuk Eagle Sand»", ""),
        ("Piezas (88)", "Acabado exacto del showroom", "Se guarda con la pieza; el asistente todavía no lo menciona", "FEB-001: «Nero Marquinia»", ""),
        ("Piezas (88) y Modelos", "Las columnas «(lo que entiende el asistente)»", "Corrige cómo lo clasifica; se usa en lugar de lo que dedujo del texto", "Navigli: forma → «modular» (sin «puertas corredizas»)", ""),
        ("Paletas por modelo", "Decisión", "aplica = la ofrece bajo pedido · aviso = la menciona pidiendo confirmar en la ficha · no aplica = no la menciona", "Couple · SPECCHIO → no aplica", ""),
        ("Sinónimos", "Nombres y sinónimos", "Las palabras que entiende, en español, italiano e inglés", "sofá: agregar «love seat»", ""),
        ("Colores → familia", "Familia(s)", "Con qué color aparece al buscar; vale para todas las opciones con ese nombre", "«Moss» → verde", ""),
        ("Materiales por colección", "Material", "Con qué material aparece la colección", "LACCATO LUCIDO → lacado brillante", ""),
        ("Armonías", "Distancia y estado", "Qué combina o se parece; «borrador» = solo sugerencia, «aprobado» = criterio de la marca", "amarillo combina con café → aprobado", ""),
        ("Ánimos", "Lo que prefiere mostrar", "Qué muestra primero cuando piden algo «acogedor»; ordena, nunca filtra", "acogedor → terciopelo, madera", ""),
    ])

    R.title("Hojas de consulta (no se editan)")
    R.table(["Hoja", "Qué es", "De dónde sale", "", ""], [
        ("Modelos", "Lo que dice la ficha oficial de cada modelo: forma, materiales, estilo, medidas, tipo de lista", "febalcasa.com, leída a mano", "", ""),
        ("Opciones oficiales", "Todo lo que el asistente puede ofrecer bajo pedido, en el orden en que se recorre el tour", "Listas propias, texto de la ficha y paletas de línea que aplican", "", ""),
        ("Acabados de las paletas", "Lo que contienen las paletas de línea, cada acabado una vez", "febalcasa.com", "", ""),
        ("Estilos", "Qué modelos son de qué estilo", "La ficha de cada modelo", "", ""),
        ("Correcciones", "Lo que las fotos del tour corrigieron del catálogo anterior", "Fotos del tour", "", ""),
    ])

    R.title("Cómo responde con estos datos")
    R.table(["Nivel", "Cuándo", "Qué dice", "", ""], [
        ("En el showroom", "La pieza del tour es así (hoja «Piezas»)", "«Aquí lo tienes…» y un botón que lleva la cámara a la pieza", "", ""),
        ("Bajo pedido", "En el tour no está así, pero su ficha lo ofrece («Opciones oficiales»)", "«Aquí no está así, pero sí bajo pedido» y el enlace a la ficha", "", ""),
        ("De la línea, con aviso", "Solo la paleta de su línea lo tiene y no está claro que aplique («Paletas por modelo»: aviso)", "«La línea maneja…; confirma en su ficha si aplica a este modelo»", "", ""),
        ("Sin confirmar", "Nadie lo sabe todavía", "Lo dice como no confirmado y enlaza la ficha; nunca lo inventa", "", ""),
    ])

    R.title("Cómo leer algunas columnas")
    R.table(["Columna", "Significa", "", "", ""], [
        ("«… (lo que entiende el asistente)»", "Cómo interpreta el asistente el texto de al lado, en conceptos fijos (p. ej. «tortora» → beige). Si está mal, corrígelo ahí mismo", "", "", ""),
        ("«(la ficha no lo dice)»", "La web oficial no publica ese dato", "", "", ""),
        ("Material o acabado vacío en «Piezas»", "No se pudo confirmar en la foto: el asistente lo dice como no confirmado", "", "", ""),
        ("Muestra", "El color de la muestra oficial de la web", "", "", ""),
    ])

    R.title("Colores del archivo")
    R.table(["Color", "Significa", "", "", ""], [
        ("Encabezado rojo", "Columna de consulta", "", "", ""),
        ("Encabezado verde con ✎", "Columna editable (celdas verde claro)", "", "", ""),
        ("ID azul claro", "La pieza tuvo una corrección ya aplicada (ver «Correcciones»)", "", "", ""),
        ("ID naranja", "Duda abierta: está en el archivo de decisiones", "", "", ""),
        ("ID rojo", "Pendiente de arreglar en el tour", "", "", ""),
    ], fills={"Encabezado rojo": HEAD_RO, "Encabezado verde con ✎": HEAD_ED, "ID azul claro": NOTE_FILLS["CORREGIDO"],
              "ID naranja": NOTE_FILLS["DUDA"], "ID rojo": NOTE_FILLS["PENDIENTE"]})

    R.title("Resumen")
    R.para(f"Piezas del tour: {s['pieces']} · color confirmado {s['color']} · material confirmado {s['material']} · forma confirmada {s['forma']} · "
           f"acabado exacto identificado {s['finish']}")
    R.para(f"Modelos (fichas oficiales): {s['models']} · opciones que puede ofrecer bajo pedido: {s['options']}")
    R.para(f"Paletas de línea: {s['pal_rows']} colecciones por modelo · aplica {s['aplica']} · aviso {s['aviso']} · no aplica {s['no']} · "
           f"{s['pal_contents']} acabados distintos")
    R.para(f"Vocabulario: {s['concepts']} conceptos con {s['synonyms']} sinónimos · {s['colors']} nombres de color · {s['collections']} colecciones · "
           f"{s['relations']} armonías y parecidos")

    R.title("Glosario")
    R.table(["Palabra", "Significa", "", "", ""], [
        ("Finiture per NOTTE", "Acabados de la línea Dormitorio (camas, burós, cómodas, armarios); «notte» = noche", "", "", ""),
        ("GIORNO", "Línea Día: sala, comedor, librerías, aparadores", "", "", ""),
        ("Paleta de la línea", "El menú de acabados que la web repite igual en varias fichas de una línea", "", "", ""),
        ("Acabado exacto del showroom", "El acabado que tiene la pieza expuesta", "", "", ""),
        ("Gres", "Porcelánico en placa, para cubiertas y frentes", "", "", ""),
        ("Laccato opaco / lucido", "Lacado mate / brillante", "", "", ""),
        ("Similpelle", "Similpiel: no es piel", "", "", ""),
    ])


def readme_it(ws, s):
    R = Readme(ws)
    R.para("Febal Casa · database dell'assistente", Font(bold=True, size=16, color=RED))
    R.para(f"Generato il {TODAY}. È esattamente ciò che l'assistente del tour usa per rispondere.")

    R.title("Che cos'è questo file")
    R.para("Qui c'è tutto ciò che l'assistente sa dello showroom: ogni pezzo del tour, ogni modello del sito ufficiale, le finiture che può proporre, "
           "le parole che capisce e come le raggruppa. Serve a due cose: vedere con esattezza cosa sa, e correggerlo o completarlo direttamente qui. "
           "I nomi dei fogli e delle colonne sono in spagnolo; questa pagina li spiega.")

    R.title("Come correggere qualcosa")
    R.table(["Passo", "Cosa fare", "", "", ""], [
        ("1", "Trova il foglio e la riga. I filtri della prima riga aiutano.", "", "", ""),
        ("2", "Modifica solo le colonne con intestazione verde e ✎ (celle verde chiaro).", "", "", ""),
        ("3", "Scrivi come nelle altre righe: famiglie di colore e materiali con il loro nome in spagnolo (es. «verde», «lacado mate»), "
              "più valori separati da virgola. I nomi validi sono nel foglio «Sinónimos», colonna «Nombre (es)». In «Decisión» e «Estado» scegli dall'elenco.", "", "", ""),
        ("4", "Non aggiungere, eliminare o riordinare righe e non modificare le colonne con intestazione rossa: vengono dal sito ufficiale o dalle foto del tour. "
              "Se lì vedi un errore, segnalalo quando rimandi il file.", "", "", ""),
        ("5", "Salva e rimanda il file. Si controlla cosa è cambiato, si applica, si rigenera la base dell'assistente, si prova e ricevi questo file aggiornato.", "", "", ""),
    ])

    R.title("Cosa cambia nell'assistente quando modifichi")
    R.table(["Foglio", "Colonna", "Effetto", "Esempio", ""], [
        ("Piezas (88)", "Colore e materiale della parte principale, forma vista", "Come descrive il pezzo esposto e in quali ricerche compare", "FEB-036 Vivienne: materiale → «Nabuk Eagle Sand»", ""),
        ("Piezas (88)", "Acabado exacto del showroom (finitura esatta)", "Si salva con il pezzo; l'assistente non la menziona ancora", "FEB-001: «Nero Marquinia»", ""),
        ("Piezas (88) e Modelos", "Le colonne «(lo que entiende el asistente)»", "Correggono come lo classifica; si usano al posto di ciò che ha dedotto dal testo", "Navigli: forma → «modular» (senza «puertas corredizas»)", ""),
        ("Paletas por modelo", "Decisión", "aplica = la propone su ordinazione · aviso = la menziona chiedendo di verificare sulla scheda · no aplica = non la menziona", "Couple · SPECCHIO → no aplica", ""),
        ("Sinónimos", "Nomi e sinonimi", "Le parole che capisce, in spagnolo, italiano e inglese", "divano: aggiungere «love seat»", ""),
        ("Colores → familia", "Familia(s)", "Con quale colore compare nelle ricerche; vale per tutte le opzioni con quel nome", "«Moss» → verde", ""),
        ("Materiales por colección", "Material", "Con quale materiale compare la collezione", "LACCATO LUCIDO → lacado brillante", ""),
        ("Armonías", "Distanza e stato", "Cosa si abbina o si somiglia; «borrador» = solo suggerimento, «aprobado» = criterio del marchio", "giallo si abbina al marrone → aprobado", ""),
        ("Ánimos", "Cosa preferisce mostrare", "Cosa mostra per primo quando chiedono qualcosa di «accogliente»; ordina, non filtra mai", "acogedor → velluto, legno", ""),
    ])

    R.title("Fogli di sola consultazione (non si modificano)")
    R.table(["Foglio", "Che cos'è", "Da dove viene", "", ""], [
        ("Modelos", "Cosa dice la scheda ufficiale di ogni modello: forma, materiali, stile, misure, tipo di elenco", "febalcasa.com, letto a mano", "", ""),
        ("Opciones oficiales", "Tutto ciò che l'assistente può proporre su ordinazione, nell'ordine del tour", "Elenchi propri, testo della scheda e palette di linea che si applicano", "", ""),
        ("Acabados de las paletas", "Il contenuto delle palette di linea, ogni finitura una volta", "febalcasa.com", "", ""),
        ("Estilos", "Quali modelli hanno quale stile", "La scheda di ogni modello", "", ""),
        ("Correcciones", "Cosa le foto del tour hanno corretto del catalogo precedente", "Foto del tour", "", ""),
    ])

    R.title("Come risponde con questi dati")
    R.table(["Livello", "Quando", "Cosa dice", "", ""], [
        ("Nello showroom", "Il pezzo del tour è così (foglio «Piezas»)", "«Eccolo…» e un pulsante che porta la telecamera sul pezzo", "", ""),
        ("Su ordinazione", "Nel tour non è così, ma la sua scheda lo propone («Opciones oficiales»)", "«Qui non è così, ma è disponibile su ordinazione» e il link alla scheda", "", ""),
        ("Della linea, con avviso", "Solo la palette della sua linea lo ha e non è chiaro che si applichi («Paletas por modelo»: aviso)", "«La linea offre…; verifica sulla scheda se vale per questo modello»", "", ""),
        ("Non confermato", "Nessuno lo sa ancora", "Lo dice come non confermato e rimanda alla scheda; non lo inventa mai", "", ""),
    ])

    R.title("Come leggere alcune colonne")
    R.table(["Colonna", "Significa", "", "", ""], [
        ("«… (lo que entiende el asistente)»", "Come l'assistente interpreta il testo accanto, in concetti fissi (es. «tortora» → beige). Se è sbagliato, correggilo lì", "", "", ""),
        ("«(la ficha no lo dice)»", "Il sito ufficiale non pubblica quel dato", "", "", ""),
        ("Materiale o finitura vuoti in «Piezas»", "Non si è potuto confermare dalla foto: l'assistente lo dice come non confermato", "", "", ""),
        ("Muestra", "Il colore del campione ufficiale del sito", "", "", ""),
    ])

    R.title("Colori del file")
    R.table(["Colore", "Significa", "", "", ""], [
        ("Intestazione rossa", "Colonna di consultazione", "", "", ""),
        ("Intestazione verde con ✎", "Colonna modificabile (celle verde chiaro)", "", "", ""),
        ("ID azzurro", "Il pezzo ha avuto una correzione già applicata (vedi «Correcciones»)", "", "", ""),
        ("ID arancione", "Dubbio aperto: si trova nel file delle decisioni", "", "", ""),
        ("ID rosso", "Da sistemare nel tour", "", "", ""),
    ], fills={"Intestazione rossa": HEAD_RO, "Intestazione verde con ✎": HEAD_ED, "ID azzurro": NOTE_FILLS["CORREGIDO"],
              "ID arancione": NOTE_FILLS["DUDA"], "ID rosso": NOTE_FILLS["PENDIENTE"]})

    R.title("Riepilogo")
    R.para(f"Pezzi del tour: {s['pieces']} · colore confermato {s['color']} · materiale confermato {s['material']} · forma confermata {s['forma']} · "
           f"finitura esatta identificata {s['finish']}")
    R.para(f"Modelli (schede ufficiali): {s['models']} · opzioni proponibili su ordinazione: {s['options']}")
    R.para(f"Palette di linea: {s['pal_rows']} collezioni per modello · aplica {s['aplica']} · aviso {s['aviso']} · no aplica {s['no']} · "
           f"{s['pal_contents']} finiture diverse")
    R.para(f"Vocabolario: {s['concepts']} concetti con {s['synonyms']} sinonimi · {s['colors']} nomi di colore · {s['collections']} collezioni · "
           f"{s['relations']} armonie e somiglianze")

    R.title("Parole spagnole del file")
    R.table(["Spagnolo", "Italiano", "", "", ""], [
        ("Pieza · Zona · Ficha oficial", "Pezzo · Zona · Scheda ufficiale", "", "", ""),
        ("Parte principal · Otras partes", "Parte principale · Altre parti", "", "", ""),
        ("Acabado · Opción · Colección", "Finitura · Opzione · Collezione", "", "", ""),
        ("Bajo pedido", "Su ordinazione", "", "", ""),
        ("Lo que entiende el asistente", "Ciò che l'assistente capisce", "", "", ""),
        ("Decisión: aplica / aviso / no aplica", "Decisione: si applica / con avviso / non si applica", "", "", ""),
        ("Estado: borrador / aprobado", "Stato: bozza / approvato", "", "", ""),
        ("Sinónimos · Familia · Tono", "Sinonimi · Famiglia · Tono", "", "", ""),
        ("Línea Dormitorio (Finiture per NOTTE)", "Linea Notte", "", "", ""),
    ])


def main():
    d, sheets = db.build_sheets()
    wb = Workbook()
    readme = wb.active
    readme.title = "Léeme"
    leggimi = wb.create_sheet("Leggimi")
    for sh in sheets:
        write_sheet(wb, sh)
    s = stats(d, sheets)
    readme_es(readme, s)
    readme_it(leggimi, s)
    meta = wb.create_sheet("_meta")
    for k, v in (("formato", FORMAT_VERSION), ("generado", TODAY), ("catalogo", d.cat.get("data_version", ""))):
        meta.append([k, v])
    meta.sheet_state = "hidden"
    named = [f"{w.title}!{c.coordinate}" for w in wb.worksheets for row in w.iter_rows() for c in row
             if isinstance(c.value, str) and NAME_RE.search(c.value)]
    if named:
        sys.exit(f"El libro menciona personas en: {', '.join(named[:12])}")
    out = OUT
    try:
        wb.save(out)
    except PermissionError:
        out = out.replace(".xlsx", f" ({dt.datetime.now():%H%M}).xlsx")
        wb.save(out)
        print(f"AVISO: el archivo estaba abierto en Excel; se guardó como {out}")
    print(f"{out}: " + ", ".join(f"{sh.title} {len(sh.rows)}" for sh in sheets))


if __name__ == "__main__":
    main()

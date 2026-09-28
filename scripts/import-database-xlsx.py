#!/usr/bin/env python3
"""
Reads an edited copy of the database workbook (scripts/build-database-xlsx.py) and turns the
edits in its green "✎" columns into data files:

  Piezas (88)               -> product-facts/placements.json (via scripts/facts-placement.py)
  Paletas por modelo        -> product-facts/palette-decisions.json
  "lo que entiende el asistente" columns (Piezas, Modelos) -> product-facts/concept-overrides.json
  Colores → familia         -> product-facts/color-overrides.json
  Materiales por colección  -> product-facts/material-overrides.json
  Sinónimos, Armonías, Ánimos -> clients/febal-casa/ontology.overrides.json

Rows are matched by the hidden "clave" column. Changes in read-only columns, and added or
deleted rows, are reported and not imported.

    python scripts/import-database-xlsx.py <file.xlsx>            # shows what changed; writes nothing
    python scripts/import-database-xlsx.py <file.xlsx> --apply    # writes the edits, rebuilds the catalog,
                                                                  # runs the engine tests, regenerates the workbook
"""
import datetime as dt
import json
import os
import re
import shutil
import subprocess
import sys

from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8")   # the Windows console defaults to cp1252
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import febal_database as db  # noqa: E402

TODAY = dt.date.today().isoformat()
EVIDENCE = f"indicado en la base de datos del showroom ({TODAY})"
ONTOLOGY_OVERRIDES = os.path.join(db.CLIENT, "ontology.overrides.json")
COLOR_OVERRIDES = os.path.join(db.FACTS, "color-overrides.json")
MATERIAL_OVERRIDES = os.path.join(db.FACTS, "material-overrides.json")
DECISIONS_FILE = os.path.join(db.FACTS, "palette-decisions.json")
CONCEPT_OVERRIDES = os.path.join(db.FACTS, "concept-overrides.json")


def text(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).replace("\r\n", "\n").strip()


def same(a, b):
    return re.sub(r"\s+", " ", text(a)) == re.sub(r"\s+", " ", text(b))


def read_workbook(path, sheets):
    wb = load_workbook(path, data_only=True)
    out, problems = {}, []
    for sh in sheets:
        if sh.title not in wb.sheetnames:
            problems.append(f"Falta la hoja «{sh.title}»")
            continue
        ws = wb[sh.title]
        headers = [text(c.value) for c in ws[1]]
        if db.KEY not in headers:
            problems.append(f"«{sh.title}»: falta la columna oculta «{db.KEY}»; no se puede leer")
            continue
        missing = [c.header for c in sh.cols if c.header not in headers]
        if missing:
            problems.append(f"«{sh.title}»: columnas renombradas o borradas: {', '.join(missing)}")
        rows = {}
        for r in ws.iter_rows(min_row=2, values_only=True):
            row = dict(zip(headers, r))
            if text(row.get(db.KEY)):
                rows[text(row[db.KEY])] = row
        out[sh.title] = rows
    return out, problems


def diff(sheets, book):
    edits, notes = [], []
    for sh in sheets:
        rows = book.get(sh.title)
        if rows is None:
            continue
        base = {r.key: r for r in sh.rows}
        for key, r in base.items():
            if key not in rows:
                notes.append(f"[{sh.title}] {key}: la fila no está (borrada); no se importa")
                continue
            for i, c in enumerate(sh.cols):
                if c.header == db.KEY or c.header not in rows[key] or same(r.values[i], rows[key][c.header]):
                    continue
                item = (sh.title, key, c.header, text(r.values[i]), text(rows[key][c.header]))
                (edits if c.editable else notes).append(item if c.editable else
                                                        f"[{sh.title}] {key} · {c.header}: columna de consulta, no se importa («{item[3]}» → «{item[4]}»)")
        for key in rows.keys() - base.keys():
            notes.append(f"[{sh.title}] {key}: fila nueva; no se importa")
    return edits, notes


# ------------------------------------------------------------------ translate edits into data
class Plan:
    def __init__(self, d):
        self.d, self.errors = d, []
        self.pieces = {}                      # pid -> {header: new}
        self.decisions = {}                   # (model, COLLECTION) -> verdict
        self.colors, self.materials = {}, {}
        self.concepts, self.relations, self.moods = {}, {}, {}
        self.understood = {"models": {}, "pieces": {}}   # key -> {field: [ids]}

    def ids(self, value, facet, where):
        out = []
        for label in db.split_list(value):
            if facet:
                cid = self.d.by_label[facet].get(db.norm(label))
                if not cid:
                    self.errors.append(f"{where}: «{label}» no es un {db.FACET_ES.get(facet, facet)} conocido (ver «Sinónimos», columna «Nombre (es)»)")
            else:
                found = self.d.any_label.get(db.norm(label), set())
                cid = next(iter(found)) if len(found) == 1 else None
                if not cid:
                    self.errors.append(f"{where}: «{label}» " + ("no es un concepto conocido" if not found else f"es ambiguo ({', '.join(sorted(found))}); escribe el id"))
            if cid:
                out.append(cid)
        return out

    def add(self, sheet, key, header, old, new):
        where = f"[{sheet}] {key} · {header}"
        if sheet == "Piezas (88)" and header in db.PIECE_FIELDS:
            field = db.PIECE_FIELDS[header]
            if field == "material":
                # "dominante · otras partes"; "sin confirmar" = the main surface is unknown
                head, _, rest = new.partition("·")
                main = [] if db.norm(head) in ("", "sin confirmar") else self.ids(head, "material", where)
                if len(main) > 1:
                    self.errors.append(f"{where}: antes del «·» va un solo material (el de la parte principal)")
                others = self.ids(rest, "material", where)
                ids = main[:1] + [x for x in others if x not in main]
                if not main and others:
                    self.errors.append(f"{where}: escribe el material principal antes del «·» (o «sin confirmar»); por ahora solo se guarda si hay principal")
            else:
                ids = self.ids(new, db.UNDERSTOOD_FACET[field], where)
            self.understood["pieces"].setdefault(key, {})[field] = ids
        elif sheet == "Modelos" and header in db.MODEL_FIELDS:
            field = db.MODEL_FIELDS[header]
            self.understood["models"].setdefault(key, {})[field] = self.ids(new, db.UNDERSTOOD_FACET[field], where)
        elif sheet == "Piezas (88)":
            self.pieces.setdefault(key, {})[header] = new
        elif sheet == "Paletas por modelo":
            v = db.DECISION_KEYS.get(new.strip().lower())
            if not v:
                self.errors.append(f"{where}: «{new}» no vale; usa aplica, aviso o no aplica")
            self.decisions[tuple(key.split("|", 1))] = v
        elif sheet == "Colores → familia":
            self.colors[key] = self.ids(new, "color", where)
        elif sheet == "Materiales por colección":
            ids = self.ids(new, "material", where)
            if len(ids) > 1:
                self.errors.append(f"{where}: una colección lleva un solo material")
            self.materials[key] = ids[0] if ids else None
        elif sheet == "Sinónimos":
            lang = {"Nombre (es)": "es", "Nome (it)": "it", "Name (en)": "en", "Sinónimos es": "es", "Sinonimi it": "it", "Synonyms en": "en"}[header.lstrip("✎ ")]
            e = self.concepts.setdefault(key, {})
            if header.lstrip("✎ ").startswith(("Nombre", "Nome", "Name")):
                if not new.strip():
                    self.errors.append(f"{where}: el nombre no puede quedar vacío")
                e.setdefault("labels", {})[lang] = new.strip()
            else:
                e.setdefault("synonyms", {})[lang] = db.split_list(new)
        elif sheet == "Armonías":
            e = self.relations.setdefault(key, {})
            if header.startswith("✎ Distancia"):
                try:
                    x = float(str(new).replace(",", "."))
                    assert 0 <= x <= 1
                    e["distance"] = x
                except (ValueError, AssertionError):
                    self.errors.append(f"{where}: la distancia es un número de 0 a 1")
            else:
                s = db.STATUS_KEYS.get(new.strip().lower())
                if not s:
                    self.errors.append(f"{where}: «{new}» no vale; usa borrador o aprobado")
                e["status"] = s
        elif sheet == "Ánimos":
            self.moods[key] = self.ids(new, None, where)

    # ---- writers
    def write(self):
        written = []
        if self.pieces:
            self._write_pieces()
            written.append("placements.json")
        if self.decisions:
            self._write_decisions()
            written.append("palette-decisions.json")
        if self.colors:
            merge_json(COLOR_OVERRIDES, self.colors)
            written.append("color-overrides.json")
        if self.materials:
            merge_json(MATERIAL_OVERRIDES, self.materials)
            written.append("material-overrides.json")
        if self.understood["models"] or self.understood["pieces"]:
            o = db.load_json(CONCEPT_OVERRIDES, {})
            for part in ("models", "pieces"):
                for key, fields in self.understood[part].items():
                    o.setdefault(part, {}).setdefault(key, {}).update(fields)
            with open(CONCEPT_OVERRIDES, "w", encoding="utf-8") as f:
                json.dump(o, f, ensure_ascii=False, indent=1)
            written.append("concept-overrides.json")
        if self.concepts or self.relations or self.moods:
            self._write_ontology()
            written.append("ontology.overrides.json")
        return written

    def _write_pieces(self):
        L = lambda b: [b["confirmado"], b["valor"], b["evidencia"]]
        records = []
        for pid, changes in self.pieces.items():
            p = self.d.placements[pid]
            comps = [dict(c) for c in p["componentes"]] or [{"rol": "pieza", "dominante": True, "color": "", "material": ""}]
            dom = next((c for c in comps if c.get("dominante")), comps[0])
            rec = {"product_id": pid, "visible": p["visible"], "componentes": comps, "color": L(p["color"]), "material": L(p["material"]),
                   "forma": L(p["forma"]), "opcion_oficial": L(p["opcion_oficial"]), "notas": list(p.get("notas") or [])}
            for header, new in changes.items():
                h = header.lstrip("✎ ")
                if h == "Color de la parte principal":
                    old = dom.get("color", "")
                    dom["color"] = new
                    rec["color"] = ["SI", "; ".join(c["color"] for c in comps if c.get("color")), EVIDENCE] if new else ["NO", "", EVIDENCE]
                elif h == "Material de la parte principal":
                    old = dom.get("material", "") if p["material"]["confirmado"] == "SI" else ""
                    dom["material"] = new
                    rec["material"] = ["SI", new, EVIDENCE] if new else ["NO", "", EVIDENCE]
                elif h == "Forma vista":
                    old = p["forma"]["valor"] if p["forma"]["confirmado"] == "SI" else ""
                    rec["forma"] = ["SI", new, EVIDENCE] if new else ["NO", "", EVIDENCE]
                else:
                    old = p["opcion_oficial"]["valor"] if p["opcion_oficial"]["confirmado"] == "SI" else ""
                    rec["opcion_oficial"] = ["SI", new, EVIDENCE] if new else ["NO", "", EVIDENCE]
                rec["notas"].append(f"CORREGIDO (base de datos): {h.lower()} «{old}» → «{new}» ({TODAY})")
            records.append(rec)
        out = subprocess.run([sys.executable, os.path.join(db.ROOT, "scripts", "facts-placement.py")],
                             input=json.dumps(records, ensure_ascii=False).encode("utf-8"), capture_output=True)
        if out.returncode:
            raise SystemExit("facts-placement.py falló:\n" + out.stderr.decode("utf-8", "replace")[-800:])

    def _write_decisions(self):
        dec = db.load_json(DECISIONS_FILE, {})
        for (model, coll), v in self.decisions.items():
            e = dec.setdefault(model, {"default": "aviso", "evidencia": ""})
            for k in ("aplica", "aviso", "no_aplica"):
                if k in e:
                    e[k] = [x for x in e[k] if x.strip().upper() != coll]
                    if not e[k]:
                        del e[k]
            if v != e["default"]:
                e.setdefault(v, []).append(coll)
            e["editado"] = TODAY
        with open(DECISIONS_FILE, "w", encoding="utf-8") as f:
            f.write(dump_compact(dec))

    def _write_ontology(self):
        o = db.load_json(ONTOLOGY_OVERRIDES, {})
        for cid, e in self.concepts.items():
            cur = o.setdefault("concepts", {}).setdefault(cid, {})
            for part in ("labels", "synonyms"):
                if part in e:
                    cur.setdefault(part, {}).update(e[part])
        rels = o.setdefault("relations", []) if self.relations else o.get("relations", [])
        for key, e in self.relations.items():
            typ, frm, to = key.split("|")
            cur = next((r for r in rels if (r["type"], r["from"], r["to"]) == (typ, frm, to)), None)
            if not cur:
                cur = {"from": frm, "to": to, "type": typ}
                rels.append(cur)
            cur.update(e)
        if self.moods:
            o.setdefault("moods", {}).update(self.moods)
        with open(ONTOLOGY_OVERRIDES, "w", encoding="utf-8") as f:
            json.dump(o, f, ensure_ascii=False, indent=1)


def merge_json(path, updates):
    data = db.load_json(path, {})
    data.update(updates)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(dict(sorted(data.items())), f, ensure_ascii=False, indent=1)


def dump_compact(dec):
    """palette-decisions.json keeps one line per field, lists inline."""
    j = lambda v: json.dumps(v, ensure_ascii=False)
    out = []
    for k, v in dec.items():
        if isinstance(v, dict):
            out.append(f" {j(k)}: {{\n" + ",\n".join(f"  {j(kk)}: {j(vv)}" for kk, vv in v.items()) + "\n }")
        else:
            out.append(f" {j(k)}: {j(v)}")
    return "{\n" + ",\n".join(out) + "\n}\n"


def run(cmd, cwd=db.ROOT):
    print(f"\n$ {cmd}")
    out = subprocess.run(cmd, cwd=cwd, shell=True, capture_output=True)
    tail = (out.stdout + out.stderr).decode("utf-8", "replace").strip().splitlines()
    print("\n".join(tail[-12:]))
    if out.returncode:
        raise SystemExit(f"Falló: {cmd}")


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    path, apply = sys.argv[1], "--apply" in sys.argv
    d, sheets = db.build_sheets()
    book, problems = read_workbook(path, sheets)
    edits, notes = diff(sheets, book)
    plan = Plan(d)
    for e in edits:
        plan.add(*e)

    print(f"Archivo: {path}")
    print(f"Cambios para importar: {len(edits)}")
    for sheet, key, header, old, new in edits:
        print(f"  [{sheet}] {key} · {header.lstrip('✎ ')}: «{old}» → «{new}»")
    for title, items in (("No se importa", problems + notes), ("Errores (corregir en el Excel antes de aplicar)", plan.errors)):
        if items:
            print(f"{title}: {len(items)}")
            for x in items:
                print(f"  {x}")
    if plan.errors:
        sys.exit(1)
    if not apply or not edits:
        if edits:
            print("\nNada se escribió. Para aplicar: agrega --apply")
        return

    os.makedirs(os.path.join(db.ROOT, ".scratch", "database-imports"), exist_ok=True)
    shutil.copy(path, os.path.join(db.ROOT, ".scratch", "database-imports", f"{dt.datetime.now():%Y%m%d-%H%M%S}-{os.path.basename(path)}"))
    print("\nEscrito: " + ", ".join(plan.write()))
    run("npx tsx scripts/build-catalog-v2.ts febal-casa --strict")
    run("npm test", cwd=os.path.join(db.ROOT, "packages", "assistant-engine"))
    run(f'"{sys.executable}" scripts/build-database-xlsx.py')


if __name__ == "__main__":
    main()
